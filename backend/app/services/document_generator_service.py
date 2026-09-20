"""
Document Generator service — generates DOCX files for legal cases.
"""

from __future__ import annotations

import io
import json
import uuid
from typing import Any

from docx import Document
from sqlalchemy.ext.asyncio import AsyncSession

from app.config.contacts import BENEFICIAL_CONTACTS
from app.config.settings import settings
from app.integrations.vertex_ai_client import VertexAIClient, get_vertex_ai_client
from app.prompts.document_generator import DOCUMENT_DRAFTER
from app.repositories.case_file_repository import CaseFileRepository
from app.utils.logger import get_logger
from app.services.model_config_service import cheap_model, strong_model

logger = get_logger(__name__)


class DocumentGeneratorService:
    """Generates official legal documents as DOCX files."""

    def __init__(self, gemini_client: VertexAIClient | None = None) -> None:
        self._gemini = gemini_client

    @property
    def gemini(self) -> VertexAIClient:
        if self._gemini is None:
            self._gemini = get_vertex_ai_client()
        return self._gemini

    async def generate_document(
        self,
        db: AsyncSession,
        case_file_id: uuid.UUID,
        document_type: str = "ოფიციალური დოკუმენტი / სარჩელი / საჩივარი",
        notes_for_drafting: str | None = None,
    ) -> dict[str, Any]:
        """
        Drafts a legal document based on a CaseFile.
        Returns a dict with 'docx_bytes', 'dispatch_info', and 'markdown_text'.
        """
        repo = CaseFileRepository(db)
        cf = await repo.get_by_id(case_file_id)
        if not cf:
            raise ValueError("Case file not found")

        # Prepare case data summary
        case_data = {
            "title": cf.title,
            "facts": cf.facts,
            "evidence": cf.evidence,
            "applicable_laws": cf.applicable_laws,
            "defense_strategies": cf.defense_strategies,
        }
        if notes_for_drafting:
            case_data["user_notes_for_drafting"] = notes_for_drafting

        # Prepare contacts
        contacts_data = [c.model_dump() for c in BENEFICIAL_CONTACTS]

        prompt = DOCUMENT_DRAFTER.render(
            document_type=document_type,
            case_data=json.dumps(case_data, ensure_ascii=False, indent=2),
            contacts_data=json.dumps(contacts_data, ensure_ascii=False, indent=2),
        )

        logger.info("document_generation_start", case_id=str(case_file_id))

        result = await self.gemini.generate_json(
            prompt=prompt,
            temperature=DOCUMENT_DRAFTER.temperature,
            model_name=await strong_model(),
        )

        if not isinstance(result, dict) or "document_markdown" not in result:
            raise ValueError("Failed to generate document text from AI")

        markdown_text = result["document_markdown"]
        dispatch_info = result.get("dispatch_info", {})

        # Create DOCX
        docx_bytes = self._create_docx_from_markdown(markdown_text)

        logger.info("document_generation_success", case_id=str(case_file_id))

        return {
            "docx_bytes": docx_bytes,
            "dispatch_info": dispatch_info,
            "markdown_text": markdown_text,
        }

    def _create_docx_from_markdown(self, markdown_text: str) -> bytes:
        """Converts simple markdown text to a DOCX byte array."""
        doc = Document()
        
        # We will parse very basic markdown: lines starting with #, **, etc.
        lines = markdown_text.split("\n")
        
        for line in lines:
            line = line.strip()
            if not line:
                doc.add_paragraph("")
                continue
                
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            else:
                p = doc.add_paragraph()
                # Basic bold parsing: **text**
                parts = line.split("**")
                for i, part in enumerate(parts):
                    if i % 2 == 1:
                        # Bold text
                        run = p.add_run(part)
                        run.bold = True
                    else:
                        p.add_run(part)
                        
        # Save to bytes buffer
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.read()


_doc_gen_service: DocumentGeneratorService | None = None

def get_document_generator_service() -> DocumentGeneratorService:
    global _doc_gen_service
    if _doc_gen_service is None:
        _doc_gen_service = DocumentGeneratorService()
    return _doc_gen_service
