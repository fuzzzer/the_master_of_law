"""
Tests for document generator service — DOCX creation from markdown.
"""

from __future__ import annotations

import io
from docx import Document

import pytest

from app.services.document_generator_service import DocumentGeneratorService


class TestCreateDocxFromMarkdown:
    """Test _create_docx_from_markdown method."""

    def setup_method(self):
        self.svc = DocumentGeneratorService.__new__(DocumentGeneratorService)
        self.svc._gemini = None

    def test_headings(self):
        """Markdown headings should become DOCX headings."""
        markdown = "# Title\n## Subtitle\n### Sub-subtitle"
        result = self.svc._create_docx_from_markdown(markdown)
        assert isinstance(result, bytes)
        assert len(result) > 0

        # Parse the DOCX to verify headings
        doc = Document(io.BytesIO(result))
        paragraphs = doc.paragraphs
        # Find heading paragraphs
        heading_texts = [p.text for p in paragraphs if p.style.name.startswith("Heading")]
        assert "Title" in heading_texts
        assert "Subtitle" in heading_texts
        assert "Sub-subtitle" in heading_texts

    def test_bold_text(self):
        """**bold text** should produce a bold run."""
        markdown = "This is **bold text** in a sentence."
        result = self.svc._create_docx_from_markdown(markdown)
        assert isinstance(result, bytes)

        doc = Document(io.BytesIO(result))
        # Find the paragraph with the bold text
        for p in doc.paragraphs:
            for run in p.runs:
                if run.text == "bold text":
                    assert run.bold is True

    def test_empty_lines(self):
        """Empty lines should produce empty paragraphs."""
        markdown = "Line 1\n\nLine 2"
        result = self.svc._create_docx_from_markdown(markdown)
        assert isinstance(result, bytes)

        doc = Document(io.BytesIO(result))
        # Should have at least 3 paragraphs (Line 1, empty, Line 2)
        assert len(doc.paragraphs) >= 3

    def test_output_is_bytes(self):
        """Any markdown should return bytes with length > 0."""
        markdown = "Hello world"
        result = self.svc._create_docx_from_markdown(markdown)
        assert isinstance(result, bytes)
        assert len(result) > 0

    def test_complex_document(self):
        """Full document with multiple element types."""
        markdown = (
            "# Document Title\n"
            "\n"
            "## Section 1\n"
            "Regular paragraph text.\n"
            "**Important** information here.\n"
            "\n"
            "### Subsection\n"
            "More details about the case.\n"
        )
        result = self.svc._create_docx_from_markdown(markdown)
        assert isinstance(result, bytes)
        assert len(result) > 0

        doc = Document(io.BytesIO(result))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Document Title" in all_text
        assert "Section 1" in all_text
        assert "Regular paragraph text" in all_text
